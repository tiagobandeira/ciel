"""Cria um documento Word (.docx) a partir de texto com marcação simples."""

REQUIREMENTS = ['python-docx']

import re

from docx import Document


def _com_negrito(p, texto: str) -> None:
    for parte in re.split(r'(\*\*.+?\*\*)', texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**') and len(parte) > 4:
            r = p.add_run(parte[2:-2])
            r.bold = True
        else:
            p.add_run(parte)


def run(caminho: str, conteudo: str) -> str:
    """
    caminho: caminho do arquivo .docx de saída
    conteudo: texto com marcação simples — '# ' título, '## ' subtítulo, '- ' item de lista, '> ' citação, '**negrito**' inline; linhas em branco são ignoradas
    """
    try:
        if not caminho.lower().endswith('.docx'):
            return 'Erro: o caminho de saída deve terminar em .docx'
        doc = Document()
        for bruta in conteudo.splitlines():
            linha = bruta.strip()
            if not linha:
                continue
            if linha.startswith('## '):
                doc.add_heading(linha[3:].strip(), level=2)
            elif linha.startswith('# '):
                doc.add_heading(linha[2:].strip(), level=1)
            elif linha.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                _com_negrito(p, linha[2:].strip())
            elif linha.startswith('> '):
                try:
                    p = doc.add_paragraph(style='Intense Quote')
                except KeyError:
                    p = doc.add_paragraph()
                _com_negrito(p, linha[2:].strip())
            else:
                p = doc.add_paragraph()
                _com_negrito(p, linha)
        doc.save(caminho)
        return f'Documento criado com sucesso: {caminho}'
    except Exception as e:
        return f'Erro ao criar documento: {e}'
